"""
tests/test_document_ingestion.py

Unit tests for PDF and DOCX document ingestion pipeline.
Tests validation, extraction, cleaning, table parsing, technical identifier preservation,
metadata creation, chunking, deduplication, and ChromaDB insertion in isolated environments.
"""

import os
import shutil
import tempfile
import unittest
from pathlib import Path

import docx
import pymupdf
import chromadb

from src.document_ingestion import (
    validate_document_file,
    clean_text,
    ingest_document,
    SUPPORTED_EXTENSIONS,
    IngestionError,
)

PROJECT_ROOT = Path(__file__).resolve().parent.parent


class TestDocumentIngestion(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.chroma_temp_dir = Path(self.temp_dir) / "test_chroma"

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    # --------------------------------------------------------------------------
    # 1. Validation Tests
    # --------------------------------------------------------------------------
    def test_supported_extensions(self):
        self.assertIn(".pdf", SUPPORTED_EXTENSIONS)
        self.assertIn(".docx", SUPPORTED_EXTENSIONS)

    def test_unsupported_extension_rejected(self):
        fake_txt = Path(self.temp_dir) / "manual.txt"
        fake_txt.write_text("Some text", encoding="utf-8")
        with self.assertRaises(ValueError) as ctx:
            validate_document_file(fake_txt)
        self.assertIn("Unsupported document format '.txt'", str(ctx.exception))

        fake_jpg = Path(self.temp_dir) / "photo.jpg"
        fake_jpg.write_bytes(b"\xff\xd8\xff")
        with self.assertRaises(ValueError) as ctx2:
            validate_document_file(fake_jpg)
        self.assertIn("Unsupported document format '.jpg'", str(ctx2.exception))

    def test_missing_file_raises_filenotfound(self):
        non_existent = Path(self.temp_dir) / "does_not_exist.pdf"
        with self.assertRaises(FileNotFoundError):
            validate_document_file(non_existent)

    def test_empty_file_raises_valueerror(self):
        empty_pdf = Path(self.temp_dir) / "empty.pdf"
        empty_pdf.touch()
        with self.assertRaises(ValueError) as ctx:
            validate_document_file(empty_pdf)
        self.assertIn("is empty", str(ctx.exception))

    # --------------------------------------------------------------------------
    # 2. Text Cleaning & Technical Identifiers
    # --------------------------------------------------------------------------
    def test_technical_identifiers_survive_cleaning(self):
        raw_text = (
            "Vantor Industrial (fictional) - SYNTHETIC SAMPLE DOCUMENTATION - not for real equipment Page 2\n"
            "   The WP-400 machine raised alarm E-102. Working pressure was 6.0 bar (0.6 MPa).\n"
            "Replace filter cartridge WP4-FL-DE1 or belt WP4-BT-85 when delta P exceeds 1200 Pa.\n"
            "Use grease LG-2 (WP4-GR-LG2). For WR-700 use collet WR7-CL-12 and alarm R-101.\n"
            "Fiber laser error L-05 requires replacement window MF1-PW-27.\n"
        )
        cleaned = clean_text(raw_text)

        # Technical identifiers must remain intact
        self.assertIn("WP-400", cleaned)
        self.assertIn("E-102", cleaned)
        self.assertIn("6.0 bar", cleaned)
        self.assertIn("WP4-FL-DE1", cleaned)
        self.assertIn("WP4-BT-85", cleaned)
        self.assertIn("1200 Pa", cleaned)
        self.assertIn("LG-2", cleaned)
        self.assertIn("WR-700", cleaned)
        self.assertIn("WR7-CL-12", cleaned)
        self.assertIn("R-101", cleaned)
        self.assertIn("L-05", cleaned)
        self.assertIn("MF1-PW-27", cleaned)

        # Header/footer noise should be stripped
        self.assertNotIn("SYNTHETIC SAMPLE DOCUMENTATION", cleaned)

    def test_clean_text_normalizes_whitespace_and_newlines(self):
        raw = "Line 1\r\n\r\n\r\n\r\nLine 2   with   spaces\xa0and non-breaking space"
        cleaned = clean_text(raw)
        self.assertNotIn("\r", cleaned)
        self.assertNotIn("\xa0", cleaned)
        self.assertNotIn("\n\n\n", cleaned)
        self.assertIn("Line 1\n\nLine 2 with spaces and non-breaking space", cleaned)

    # --------------------------------------------------------------------------
    # 3. PDF Extraction & Ingestion
    # --------------------------------------------------------------------------
    def test_pdf_extraction_and_ingestion(self):
        # Create a small valid test PDF
        test_pdf_path = Path(self.temp_dir) / "test_manual.pdf"
        doc = pymupdf.open()
        p1 = doc.new_page()
        p1.insert_text((50, 50), "Overview Page\nThis is page 1 introduction.")
        p2 = doc.new_page()
        p2.insert_text((50, 50), "WP-400 Saw Operation Guide\nEnsure working pressure is 6.0 bar before start.")
        doc.save(str(test_pdf_path))
        doc.close()

        res = ingest_document(
            file_path=test_pdf_path,
            collection_name="test_col_pdf",
            chroma_dir=self.chroma_temp_dir
        )

        self.assertEqual(res["file_type"], "pdf")
        self.assertEqual(res["source_file"], "test_manual.pdf")
        self.assertEqual(res["pages"], 2)
        self.assertGreaterEqual(res["chunks_created"], 2)
        self.assertEqual(res["chunks_inserted"], res["chunks_created"])

        # Check in ChromaDB
        client = chromadb.PersistentClient(path=str(self.chroma_temp_dir))
        col = client.get_collection("test_col_pdf")
        self.assertEqual(col.count(), res["chunks_inserted"])

        records = col.get()
        pages = [m["page"] for m in records["metadatas"]]
        self.assertIn(1, pages)
        self.assertIn(2, pages)
        file_types = [m["file_type"] for m in records["metadatas"]]
        self.assertTrue(all(ft == "pdf" for ft in file_types))

    # --------------------------------------------------------------------------
    # 4. DOCX Extraction, Paragraphs & Tables
    # --------------------------------------------------------------------------
    def test_docx_paragraph_and_table_extraction(self):
        test_docx_path = Path(self.temp_dir) / "test_guide.docx"
        doc = docx.Document()
        doc.add_heading("WR-700 CNC Router Guide", level=1)
        doc.add_paragraph("Daily Startup Procedure")
        doc.add_paragraph("Check compressed air gauge at 7.0 bar. Never operate under 6.0 bar.")

        # Add a table (like parts catalog)
        table = doc.add_table(rows=1, cols=4)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Part No."
        hdr_cells[1].text = "Description"
        hdr_cells[2].text = "Fits"
        hdr_cells[3].text = "Replace when"

        row_cells = table.add_row().cells
        row_cells[0].text = "WR7-CL-12"
        row_cells[1].text = "ER32 collet, 12 mm"
        row_cells[2].text = "WR-700"
        row_cells[3].text = "Every 800 hours"

        doc.save(str(test_docx_path))

        res = ingest_document(
            file_path=test_docx_path,
            collection_name="test_col_docx",
            chroma_dir=self.chroma_temp_dir
        )

        self.assertEqual(res["file_type"], "docx")
        self.assertEqual(res["source_file"], "test_guide.docx")
        self.assertEqual(res["tables"], 1)
        self.assertGreaterEqual(res["chunks_created"], 1)

        # Verify in ChromaDB that table text is stored and searchable
        client = chromadb.PersistentClient(path=str(self.chroma_temp_dir))
        col = client.get_collection("test_col_docx")
        records = col.get()
        stored_texts = " ".join(records["documents"])

        # Table content and technical identifiers must be in ChromaDB
        self.assertIn("WR7-CL-12", stored_texts)
        self.assertIn("ER32 collet", stored_texts)
        self.assertIn("Table:", stored_texts)
        self.assertIn("7.0 bar", stored_texts)

        # Metadata checks
        self.assertEqual(records["metadatas"][0]["file_type"], "docx")
        self.assertEqual(records["metadatas"][0]["source_file"], "test_guide.docx")

    # --------------------------------------------------------------------------
    # 5. Deduplication Safety
    # --------------------------------------------------------------------------
    def test_duplicate_ingestion_does_not_create_duplicate_chunks(self):
        test_docx_path = Path(self.temp_dir) / "single_file.docx"
        doc = docx.Document()
        doc.add_heading("Section 1", level=1)
        doc.add_paragraph("Paragraph content for testing deduplication.")
        doc.save(str(test_docx_path))

        # First ingestion
        res1 = ingest_document(test_docx_path, collection_name="test_dedup", chroma_dir=self.chroma_temp_dir)
        initial_chunks = res1["chunks_inserted"]

        client = chromadb.PersistentClient(path=str(self.chroma_temp_dir))
        col = client.get_collection("test_dedup")
        self.assertEqual(col.count(), initial_chunks)

        # Second ingestion of the same file
        res2 = ingest_document(test_docx_path, collection_name="test_dedup", chroma_dir=self.chroma_temp_dir)
        self.assertEqual(res2["chunks_inserted"], initial_chunks)
        self.assertEqual(col.count(), initial_chunks, "Chunk count must remain identical after re-ingesting the same document.")


if __name__ == "__main__":
    unittest.main()
