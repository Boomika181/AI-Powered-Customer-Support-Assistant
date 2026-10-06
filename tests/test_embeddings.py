"""
tests/test_embeddings.py

Unit tests for Gemini Embedding 2 module (src/embeddings.py) and RAG integration.
Uses offline mocks to validate:
1. Correct model name (gemini-embedding-2)
2. Correct dimensionality (768)
3. Document embedding generation
4. Query embedding generation
5. PDF ingestion with Gemini Embedding 2
6. DOCX ingestion with Gemini Embedding 2
7. Metadata preservation in ChromaDB
8. ChromaDB insertion with 768 dimensions
9. ChromaDB retrieval with query vector
10. Strict failure without MiniLM fallback (preserves original exception)
11. Dimension mismatch protection
"""

import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch, MagicMock

import docx
import pymupdf
import chromadb

from src.embeddings import (
    embed_texts,
    embed_text,
    embed_query,
    embed_documents,
    get_embedding_model_name,
    get_embedding_dimension,
    GeminiEmbeddingFunction,
    EmbeddingError,
    DimensionMismatchError,
    DEFAULT_EMBEDDING_MODEL,
    DEFAULT_EMBEDDING_DIMENSION,
)
from src.document_ingestion import ingest_document
from src.rag_search import search_knowledge_base

PROJECT_ROOT = Path(__file__).resolve().parent.parent


def _create_mock_response(num_items: int, dim: int = 768):
    """Helper to create a mock EmbedContentResponse."""
    mock_res = MagicMock()
    embeddings = []
    for _ in range(num_items):
        mock_emb = MagicMock()
        mock_emb.values = [0.01 * (i + 1) for i in range(dim)]
        embeddings.append(mock_emb)
    mock_res.embeddings = embeddings
    return mock_res


class TestGeminiEmbedding2(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.chroma_dir = Path(self.temp_dir) / "test_chroma"

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    # --------------------------------------------------------------------------
    # 1. Correct Model Name
    # --------------------------------------------------------------------------
    def test_correct_model_name_default_and_env(self):
        with patch.dict(os.environ, {}, clear=False):
            os.environ.pop("GEMINI_EMBEDDING_MODEL", None)
            self.assertEqual(get_embedding_model_name(), "gemini-embedding-2")

        with patch.dict(os.environ, {"GEMINI_EMBEDDING_MODEL": "gemini-embedding-2-custom"}):
            self.assertEqual(get_embedding_model_name(), "gemini-embedding-2-custom")

    # --------------------------------------------------------------------------
    # 2. Correct Dimensionality
    # --------------------------------------------------------------------------
    def test_correct_dimensionality_default(self):
        self.assertEqual(get_embedding_dimension(), 768)
        self.assertEqual(DEFAULT_EMBEDDING_DIMENSION, 768)

    # --------------------------------------------------------------------------
    # 3. Document Embedding Generation
    # --------------------------------------------------------------------------
    @patch("src.embeddings.get_gemini_client")
    def test_document_embedding_generation(self, mock_get_client):
        mock_client = MagicMock()
        mock_client.models.embed_content.return_value = _create_mock_response(num_items=3, dim=768)
        mock_get_client.return_value = mock_client

        docs = ["Section 1 description", "Section 2 description", "Section 3 description"]
        embeddings = embed_documents(docs)

        self.assertEqual(len(embeddings), 3)
        self.assertEqual(len(embeddings[0]), 768)
        self.assertEqual(len(embeddings[1]), 768)
        self.assertEqual(len(embeddings[2]), 768)
        mock_client.models.embed_content.assert_called_once()
        _, kwargs = mock_client.models.embed_content.call_args
        self.assertEqual(kwargs["model"], "gemini-embedding-2")
        self.assertEqual(kwargs["config"].output_dimensionality, 768)

    # --------------------------------------------------------------------------
    # 4. Query Embedding Generation
    # --------------------------------------------------------------------------
    @patch("src.embeddings.get_gemini_client")
    def test_query_embedding_generation(self, mock_get_client):
        mock_client = MagicMock()
        mock_client.models.embed_content.return_value = _create_mock_response(num_items=1, dim=768)
        mock_get_client.return_value = mock_client

        query = "How to replace filter cartridge WP4-FL-DE1?"
        vec = embed_query(query)

        self.assertEqual(len(vec), 768)
        mock_client.models.embed_content.assert_called_once()
        _, kwargs = mock_client.models.embed_content.call_args
        self.assertEqual(kwargs["model"], "gemini-embedding-2")
        self.assertEqual(kwargs["config"].output_dimensionality, 768)

    # --------------------------------------------------------------------------
    # 5. PDF Ingestion with Gemini Embedding 2
    # --------------------------------------------------------------------------
    @patch("src.embeddings.get_gemini_client")
    def test_pdf_ingestion_with_embeddings(self, mock_get_client):
        mock_client = MagicMock()
        # Mock returns 768-dim embeddings for any batch size
        def side_effect(model, contents, config):
            return _create_mock_response(num_items=len(contents), dim=768)
        mock_client.models.embed_content.side_effect = side_effect
        mock_get_client.return_value = mock_client

        test_pdf = Path(self.temp_dir) / "manual.pdf"
        doc = pymupdf.open()
        p = doc.new_page()
        p.insert_text((50, 50), "WP-400 Panel Saw Setup\nOperating pressure 6.0 bar.")
        doc.save(str(test_pdf))
        doc.close()

        res = ingest_document(test_pdf, collection_name="test_col_pdf", chroma_dir=self.chroma_dir)
        self.assertEqual(res["file_type"], "pdf")
        self.assertEqual(res["embedding_dimension"], 768)
        self.assertEqual(res["embedding_model"], "gemini-embedding-2")

        # Verify storage in ChromaDB
        client = chromadb.PersistentClient(path=str(self.chroma_dir))
        col = client.get_collection("test_col_pdf")
        self.assertEqual(col.count(), res["chunks_inserted"])
        records = col.get(include=["embeddings", "metadatas"])
        self.assertEqual(len(records["embeddings"][0]), 768)

    # --------------------------------------------------------------------------
    # 6. DOCX Ingestion with Gemini Embedding 2
    # --------------------------------------------------------------------------
    @patch("src.embeddings.get_gemini_client")
    def test_docx_ingestion_with_embeddings(self, mock_get_client):
        mock_client = MagicMock()
        def side_effect(model, contents, config):
            return _create_mock_response(num_items=len(contents), dim=768)
        mock_client.models.embed_content.side_effect = side_effect
        mock_get_client.return_value = mock_client

        test_docx = Path(self.temp_dir) / "guide.docx"
        d = docx.Document()
        d.add_heading("WR-700 Collet Replacement", level=1)
        d.add_paragraph("Replace collet WR7-CL-12 every 800 hours.")
        d.save(str(test_docx))

        res = ingest_document(test_docx, collection_name="test_col_docx", chroma_dir=self.chroma_dir)
        self.assertEqual(res["file_type"], "docx")
        self.assertEqual(res["embedding_dimension"], 768)

        client = chromadb.PersistentClient(path=str(self.chroma_dir))
        col = client.get_collection("test_col_docx")
        self.assertEqual(col.count(), res["chunks_inserted"])
        records = col.get(include=["embeddings", "metadatas"])
        self.assertEqual(len(records["embeddings"][0]), 768)

    # --------------------------------------------------------------------------
    # 7. Metadata Preservation
    # --------------------------------------------------------------------------
    @patch("src.embeddings.get_gemini_client")
    def test_metadata_preservation(self, mock_get_client):
        mock_client = MagicMock()
        mock_client.models.embed_content.side_effect = lambda model, contents, config: _create_mock_response(len(contents), 768)
        mock_get_client.return_value = mock_client

        test_pdf = Path(self.temp_dir) / "manual_meta.pdf"
        doc = pymupdf.open()
        p = doc.new_page()
        p.insert_text((50, 50), "Section 1.1\nDoc ref: UM-WP400-1.1 | Category: Machine Operation Issues | Machine: WP-400\nPressure 6.0 bar.")
        doc.save(str(test_pdf))
        doc.close()

        ingest_document(test_pdf, collection_name="meta_col", chroma_dir=self.chroma_dir)

        client = chromadb.PersistentClient(path=str(self.chroma_dir))
        col = client.get_collection("meta_col")
        records = col.get(include=["metadatas"])
        meta = records["metadatas"][0]

        expected_keys = {"source", "source_file", "file_type", "chunk_index", "title", "section", "category", "machine", "page", "document_reference"}
        for k in expected_keys:
            self.assertIn(k, meta, f"Metadata must contain {k}")

    # --------------------------------------------------------------------------
    # 8. ChromaDB Insertion
    # --------------------------------------------------------------------------
    @patch("src.embeddings.get_gemini_client")
    def test_chromadb_insertion_dimension(self, mock_get_client):
        mock_client = MagicMock()
        mock_client.models.embed_content.side_effect = lambda model, contents, config: _create_mock_response(len(contents), 768)
        mock_get_client.return_value = mock_client

        test_docx = Path(self.temp_dir) / "insert_test.docx"
        d = docx.Document()
        d.add_heading("Safety Guidelines", level=1)
        d.add_paragraph("Lock-Out / Tag-Out procedure before opening enclosure.")
        d.save(str(test_docx))

        ingest_document(test_docx, collection_name="insert_col", chroma_dir=self.chroma_dir)

        client = chromadb.PersistentClient(path=str(self.chroma_dir))
        col = client.get_collection("insert_col")
        sample = col.get(limit=1, include=["embeddings"])
        self.assertEqual(len(sample["embeddings"][0]), 768)

    # --------------------------------------------------------------------------
    # 9. ChromaDB Retrieval with Query Vector
    # --------------------------------------------------------------------------
    @patch("src.embeddings.get_gemini_client")
    def test_chromadb_retrieval(self, mock_get_client):
        mock_client = MagicMock()
        mock_client.models.embed_content.side_effect = lambda model, contents, config: _create_mock_response(len(contents), 768)
        mock_get_client.return_value = mock_client

        test_docx = Path(self.temp_dir) / "retrieval_test.docx"
        d = docx.Document()
        d.add_heading("Laser Alarm L-05", level=1)
        d.add_paragraph("Alarm L-05 indicates protective window contamination on fiber laser.")
        d.save(str(test_docx))

        ingest_document(test_docx, collection_name="retrieval_col", chroma_dir=self.chroma_dir)

        results = search_knowledge_base(
            query="laser protective window alarm L-05",
            top_k=1,
            collection_name="retrieval_col",
            chroma_dir=self.chroma_dir
        )
        self.assertEqual(len(results), 1)
        self.assertIn("L-05", results[0]["text"])
        self.assertGreater(results[0]["similarity_score"], 0.0)
        self.assertIn("embedding_latency_ms", results[0])
        self.assertIn("retrieval_latency_ms", results[0])

    # --------------------------------------------------------------------------
    # 10. Strict Failure Without MiniLM Fallback
    # --------------------------------------------------------------------------
    @patch("src.embeddings.get_gemini_client")
    def test_no_minilm_fallback_on_failure(self, mock_get_client):
        mock_client = MagicMock()
        mock_client.models.embed_content.side_effect = RuntimeError("Quota exhausted 429")
        mock_get_client.return_value = mock_client

        # Must raise EmbeddingError and preserve original exception
        with self.assertRaises(EmbeddingError) as ctx:
            embed_texts(["Technical document chunk text"])

        self.assertIn("Quota exhausted 429", str(ctx.exception))
        self.assertIsInstance(ctx.exception.__cause__, RuntimeError)

    # --------------------------------------------------------------------------
    # 11. Dimension Mismatch Protection
    # --------------------------------------------------------------------------
    @patch("src.embeddings.get_gemini_client")
    def test_dimension_mismatch_protection(self, mock_get_client):
        mock_client = MagicMock()
        # Returns 384 dimensions when 768 was expected
        mock_client.models.embed_content.return_value = _create_mock_response(num_items=1, dim=384)
        mock_get_client.return_value = mock_client

        with self.assertRaises(DimensionMismatchError) as ctx:
            embed_text("Some text", output_dimensionality=768)

        self.assertIn("expected 768, but model 'gemini-embedding-2' returned vector of length 384", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
